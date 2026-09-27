import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
p='out/CE323_Lab_Dataset.xlsx'; wb=openpyxl.load_workbook(p)
hdr=PatternFill('solid',fgColor='1F3A5F'); sev={'HIGH':'F8CBAD','MEDIUM':'FFE699','LOW':'DDEBF7','INFO':'EDEDED'}
for ws in wb.worksheets:
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=hdr; c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2'
    for i,col in enumerate(ws.iter_cols(min_row=1,max_row=min(ws.max_row,200)),1):
        L=max((len(str(c.value)) for c in col if c.value is not None),default=8)
        ws.column_dimensions[get_column_letter(i)].width=min(max(10,L+2),70)
    if ws.max_row>1: ws.auto_filter.ref=ws.dimensions
wb['README'].column_dimensions['A'].width=160
for ws in (wb['Validation_Log'],wb['Catalog']):
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
v=wb['Validation_Log']
for row in v.iter_rows(min_row=2):
    row[1].fill=PatternFill('solid',fgColor=sev.get(row[1].value,'FFFFFF'))
r=wb['Reported_Results']; ci=[c.value for c in r[1]].index('check')
for row in r.iter_rows(min_row=2):
    if row[ci].value=='MISMATCH': row[ci].fill=PatternFill('solid',fgColor='F8CBAD')
    elif row[ci].value=='MATCH': row[ci].fill=PatternFill('solid',fgColor='C6EFCE')
wb.save(p); print('styled')
